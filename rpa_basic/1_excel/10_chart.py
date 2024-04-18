from openpyxl import load_workbook
wb = load_workbook("Sample2.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference, LineChart
#B2:C11까지의 데이터 차트 생성
#bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
#bar_chart = BarChart() #차트 종류
#bar_chart.add_data(bar_value) #차트 데이터 추가
#ws.add_chart(bar_chart, "E1") # 차트를 넣을 위치



line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart() #차트 종류
line_chart.add_data(line_value, titles_from_data=True)
line_chart.title = "성적표"
line_chart.style = 20
line_chart.y_axis.title ="점수"
line_chart.x_axis.title ="번호"
ws.add_chart(line_chart, "E1") # 차트를 넣을 위치


wb.save("sample_Lchart.xlsx")


